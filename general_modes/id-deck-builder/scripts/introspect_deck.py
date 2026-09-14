"""Extract the house dialect from an example deck before building anything.

    python introspect_deck.py C:/path/example.pptx            # layouts + theme + per-slide summary
    python introspect_deck.py C:/path/example.pptx --slide 6  # every shape on one slide, with runs

Prints: canvas size, theme fonts and colour scheme, every layout name (what
Deck(layouts=...) can map to) with its placeholder count, then per slide the layout
used and a shape inventory (type, position in inches, fill, first text run's font /
size / colour). Bounded: one line per shape, text truncated - read it, do not guess.
"""
import argparse
import sys

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu

sys.stdout.reconfigure(encoding="utf-8", errors="replace")


def inches(emu):
    return round(Emu(emu).inches, 2)


def fill_of(shape):
    try:
        if shape.fill.type == 1:  # solid
            return str(shape.fill.fore_color.rgb)
    except Exception:
        pass
    return "-"


def first_run(shape):
    if not shape.has_text_frame:
        return ""
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            if r.text.strip():
                f = r.font
                col = "-"
                try:
                    col = str(f.color.rgb) if f.color and f.color.type else "-"
                except Exception:
                    pass
                size = f.size.pt if f.size else "-"
                return f"{f.name or '-'} {size} {col} {'B' if f.bold else ''} | {r.text.strip()[:50]!r}"
    return ""


def theme_summary(prs):
    master = prs.slide_masters[0]
    theme_part = None
    for rel in master.part.rels.values():
        if "theme" in rel.reltype:
            theme_part = rel.target_part
            break
    if theme_part is None:
        return "theme: not found"
    root = theme_part._element if hasattr(theme_part, "_element") else None
    if root is None:
        from lxml import etree
        root = etree.fromstring(theme_part.blob)
    out = []
    fs = root.find(".//" + qn("a:fontScheme"))
    if fs is not None:
        maj = fs.find(qn("a:majorFont") + "/" + qn("a:latin"))
        mnr = fs.find(qn("a:minorFont") + "/" + qn("a:latin"))
        out.append(f"theme fonts: major={maj.get('typeface') if maj is not None else '-'} "
                   f"minor={mnr.get('typeface') if mnr is not None else '-'}")
    cs = root.find(".//" + qn("a:clrScheme"))
    if cs is not None:
        cols = []
        for child in cs:
            tag = child.tag.split("}")[1]
            v = child.find(qn("a:srgbClr"))
            s = child.find(qn("a:sysClr"))
            cols.append(f"{tag}={v.get('val') if v is not None else (s.get('lastClr') if s is not None else '?')}")
        out.append("theme colours: " + ", ".join(cols))
    return "\n".join(out)


ap = argparse.ArgumentParser()
ap.add_argument("pptx")
ap.add_argument("--slide", type=int, default=None, help="1-based slide to dump in full")
a = ap.parse_args()
if a.pptx.startswith("/"):
    sys.exit("pass a C:/... path; python-pptx cannot open Git Bash /c/... paths")

prs = Presentation(a.pptx)
print(f"canvas: {inches(prs.slide_width)} x {inches(prs.slide_height)} in; slides: {len(prs.slides)}")
print(theme_summary(prs))
print("layouts:")
for lay in prs.slide_masters[0].slide_layouts:
    print(f"  {lay.name!r}: {len(lay.placeholders)} placeholders, {len(lay.shapes)} shapes")

for i, slide in enumerate(prs.slides, 1):
    if a.slide and i != a.slide:
        continue
    print(f"\nslide {i}: layout {slide.slide_layout.name!r}, {len(slide.shapes)} shapes"
          + (", has notes" if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip() else ""))
    shapes = slide.shapes if a.slide else list(slide.shapes)[:6]
    for shp in shapes:
        print(f"  {shp.shape_type!s:22} x={inches(shp.left):5} y={inches(shp.top):5} "
              f"w={inches(shp.width):5} h={inches(shp.height):5} fill={fill_of(shp):8} {first_run(shp)}")
    if not a.slide and len(slide.shapes) > 6:
        print(f"  ... {len(slide.shapes) - 6} more (use --slide {i})")
