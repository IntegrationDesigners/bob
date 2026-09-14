"""Integration Designers (ID) house-style deck library for python-pptx.

Builds slides FROM SCRATCH on a template carcass: open an existing ID deck, drop its
slides (keeping theme, masters, layouts and logo media), then add slides with the
recipes below. Applying the template to a marp/HTML export does not work; rebuilding
on the carcass does, and python-pptx drops the orphaned media on save (a 16 MB
template gave a 10.7 MB output with no extra stripping).

Numbers were extracted from the example decks by introspection (see
introspect_deck.py), not guessed. Canvas is 10 x 5.63 in, so point sizes are small.

Usage from a build script:

    from id_deck import Deck
    d = Deck("C:/path/IBM_Bob_for_Developers_v2.pptx")   # C:/ form, never /c/
    s = d.lead("IBM Bob 2.0", "What's new", extra="Matthias Blomme")
    s = d.cards("A skill is reusable guidance", [("Focused", "One kind of work."), ...],
                takeaway="A prompt asks for an outcome. A skill captures how.")
    d.notes(s, "speaker notes text")
    d.save("C:/out/deck.pptx")
"""
import os
import re
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

# ---- house style (from the example decks) -----------------------------------
PURPLE = RGBColor(0x28, 0x1A, 0x77)
TEAL = RGBColor(0x00, 0xD7, 0xB0)
TEAL_TITLE = RGBColor(0x01, 0xD9, 0xAA)   # title-slide teal used in the ID decks
DARK = RGBColor(0x1A, 0x23, 0x32)
GRAY = RGBColor(0x4B, 0x56, 0x63)
CARD = RGBColor(0xF4, 0xF7, 0xFA)         # light card with coloured top bar
CARD2 = RGBColor(0xF2, 0xF2, 0xF4)        # tall / panel rounded card fill
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED_ON_PURPLE = RGBColor(0x8E, 0x87, 0xB5)

SANS = "IBM Plex Sans"   # matches the existing ID decks; substituted where not installed
MONO = "Consolas"        # no mono Plex on this box; a substituted mono breaks alignment

# Layout names in the ID template. Override per template via Deck(layouts={...}).
LAYOUTS = {
    "lead": "Intro / tussentitel slide",   # purple background, teal rule, logo
    "band": "Image breakdown",             # purple top band; lower half needs a white cover
    "panel": "1_Overview 2",               # left purple panel with logo
    "quote": "1_Quote no picture",         # quote furniture left, purple panel right
    "team": "1_Meet the team",             # purple full-bleed (circle trio)
    "closer": "Thank you",
}

CANVAS_W, CANVAS_H = 10.0, 5.63
CONTENT_X, CONTENT_W = 0.47, 9.06


class Deck:
    def __init__(self, template, layouts=None, drop_template_slides=True):
        if template.startswith("/"):
            raise ValueError("python-pptx cannot open Git Bash /c/... paths; pass C:/... form")
        self.prs = Presentation(template)
        self.layout_names = dict(LAYOUTS, **(layouts or {}))
        self.layouts = {lay.name: lay for lay in self.prs.slide_masters[0].slide_layouts}
        missing = [v for v in self.layout_names.values() if v not in self.layouts]
        if missing:
            raise KeyError(f"layouts not in this template: {missing}; available: {sorted(self.layouts)}")
        if drop_template_slides:
            self._drop_slides()

    # ---- carcass ------------------------------------------------------------
    def _drop_slides(self):
        sldIdLst = self.prs.slides._sldIdLst
        for sldId in list(sldIdLst):
            self.prs.part.drop_rel(sldId.get(qn("r:id")))
            sldIdLst.remove(sldId)
        # co-authoring change info from the source deck would dangle once slides differ
        for rel_id in list(self.prs.part.rels):
            rel = self.prs.part.rels[rel_id]
            if "changesInfo" in rel.reltype or "revisionInfo" in rel.reltype:
                self.prs.part.drop_rel(rel_id)

    def new_slide(self, key, keep_placeholders=False):
        slide = self.prs.slides.add_slide(self.layouts[self.layout_names[key]])
        if not keep_placeholders:
            for shp in list(slide.shapes):
                shp._element.getparent().remove(shp._element)
        return slide

    # ---- primitives ---------------------------------------------------------
    @staticmethod
    def rect(slide, x, y, w, h, fill, rounded=False, radius=None):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
            Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
        shape.line.fill.background()
        shape.shadow.inherit = False
        if rounded and radius is not None:
            shape.adjustments[0] = radius
        return shape

    @staticmethod
    def text(slide, x, y, w, h, paragraphs, size=10.5, color=DARK, bold=False, font=SANS,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=None, space_after=None):
        """paragraphs: list of paragraphs; each a string or a list of run dicts
        {t: text, b: bold, i: italic, c: color, f: font, s: size}."""
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, m, 0)
        first = True
        for para in paragraphs:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.alignment = align
            if line_spacing:
                p.line_spacing = line_spacing
            if space_after is not None:
                p.space_after = Pt(space_after)
            runs = [{"t": para}] if isinstance(para, str) else para
            for rd in runs:
                r = p.add_run()
                r.text = rd["t"]
                r.font.name = rd.get("f", font)
                r.font.size = Pt(rd.get("s", size))
                r.font.bold = rd.get("b", bold)
                r.font.italic = rd.get("i", None)
                r.font.color.rgb = rd.get("c", color)
        return box

    @staticmethod
    def oval(slide, x, y, d, fill, label=None, label_size=12, label_color=WHITE):
        o = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
        o.fill.solid()
        o.fill.fore_color.rgb = fill
        o.line.fill.background()
        o.shadow.inherit = False
        if label is not None:
            tf = o.text_frame
            tf.word_wrap = False
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = label
            r.font.name = SANS
            r.font.size = Pt(label_size)
            r.font.bold = True
            r.font.color.rgb = label_color
        return o

    @staticmethod
    def soft_shadow(shape):
        spPr = shape._element.spPr
        el = spPr.find(qn("a:effectLst"))
        if el is None:
            el = etree.SubElement(spPr, qn("a:effectLst"))
        shdw = etree.SubElement(el, qn("a:outerShdw"))
        shdw.set("blurRad", "90000")
        shdw.set("dist", "38100")
        shdw.set("dir", "5400000")
        shdw.set("rotWithShape", "0")
        clr = etree.SubElement(shdw, qn("a:srgbClr"))
        clr.set("val", "281A77")
        alpha = etree.SubElement(clr, qn("a:alpha"))
        alpha.set("val", "20000")

    @staticmethod
    def grid_x(n, gap=0.25, x0=CONTENT_X, total=CONTENT_W):
        """n equal columns across the content width: [(x, w), ...]."""
        w = (total - gap * (n - 1)) / n
        return [(x0 + i * (w + gap), w) for i in range(n)]

    @staticmethod
    def runs(text, code_color=PURPLE):
        """Split markdown-ish `code` spans into styled runs; **bold** is honoured too."""
        out = []
        for i, piece in enumerate(text.split("`")):
            if not piece:
                continue
            if i % 2:
                out.append({"t": piece, "f": MONO, "s": 10, "c": code_color})
            else:
                for j, part in enumerate(piece.split("**")):
                    if part:
                        out.append({"t": part, "b": True} if j % 2 else {"t": part})
        return out

    # ---- slide furniture ----------------------------------------------------
    def takeaway(self, slide, text_or_runs):
        """Purple bar along the bottom with a teal one-liner (the example decks'
        Rectangle-30 pattern; replaces marp p.small footnotes)."""
        runs = text_or_runs if isinstance(text_or_runs, list) else self.runs(text_or_runs)
        self.rect(slide, 0.1, 5.19, 9.8, 0.38, PURPLE)
        self.text(slide, 0.3, 5.19, 9.4, 0.38, [runs], size=10, color=TEAL,
                  anchor=MSO_ANCHOR.MIDDLE)

    def chip(self, slide, x, y, label):
        cw = 0.34 + 0.085 * len(label)
        self.rect(slide, x, y, cw, 0.34, TEAL)
        self.text(slide, x, y, cw, 0.34, [label], size=11, color=PURPLE, bold=True,
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        return cw

    def card(self, slide, x, y, w, h, heading, body, number=None, stat=None, bar=PURPLE, head_h=0.3):
        """Light card with a coloured top bar (slide 6 of the developers deck)."""
        self.rect(slide, x, y, w, h, CARD)
        self.rect(slide, x, y, w, 0.08, bar)
        tx, tw = x + 0.14, w - 0.28
        ty = y + 0.24
        if number:
            self.text(slide, tx, ty, tw, 0.42, [number], size=22, color=TEAL, bold=True)
            ty += 0.5
        if stat:
            self.text(slide, tx, ty, tw, 0.48, [stat], size=26, color=TEAL, bold=True)
            ty += 0.6
        self.text(slide, tx, ty, tw, head_h, [heading], size=13, color=PURPLE, bold=True)
        ty += head_h + 0.08
        self.text(slide, tx, ty, tw, y + h - ty - 0.1, [self.runs(body)], size=10, color=DARK,
                  line_spacing=1.1)

    def tall_card(self, slide, x, y, w, h, number, heading, body, num_size=22):
        """Tall rounded card with soft shadow, overlapping the purple band (ACE-deck pattern)."""
        c = self.rect(slide, x, y, w, h, CARD2, rounded=True, radius=0.09)
        self.soft_shadow(c)
        self.text(slide, x + 0.16, y + 0.16, w - 0.32, 0.45, [number], size=num_size, color=TEAL, bold=True)
        self.text(slide, x + 0.16, y + 0.72, w - 0.32, 0.62, [heading], size=12.5, color=PURPLE, bold=True)
        self.text(slide, x + 0.16, y + 1.4, w - 0.32, h - 1.55, [self.runs(body)], size=9.5,
                  color=GRAY, line_spacing=1.15)

    def panel_card(self, slide, x, y, w, h, number, lead, body):
        """Rounded card with a teal numbered circle on the left (panel slides)."""
        c = self.rect(slide, x, y, w, h, CARD2, rounded=True, radius=0.12)
        self.soft_shadow(c)
        self.oval(slide, x + 0.16, y + (h - 0.5) / 2, 0.5, TEAL, label=number)
        self.text(slide, x + 0.82, y + 0.16, w - 0.98, 0.3, [lead], size=10.5, color=PURPLE, bold=True)
        self.text(slide, x + 0.82, y + 0.5, w - 0.98, h - 0.62, [self.runs(body)], size=9,
                  color=GRAY, line_spacing=1.12)

    # ---- slide recipes ------------------------------------------------------
    def band(self, title, subtitle=None):
        """Content slide: purple top band with a white title, white body area below."""
        slide = self.new_slide("band")
        # cover the layout's lower-half furniture (teal dashes, label placeholders)
        self.rect(slide, 0, 1.99, CANVAS_W + 0.001, 3.65, WHITE)
        self.text(slide, CONTENT_X, 0.42, CONTENT_W, 0.5, [title], size=22, color=WHITE, bold=True)
        if subtitle:
            self.text(slide, CONTENT_X, 1.02, CONTENT_W, 0.75, [subtitle], size=10.5, color=WHITE)
        return slide

    def lead(self, title, subtitle, extra=None, title_size=36, sub_size=15, sticker=None):
        """Title or section slide: teal title, white subtitle on the purple layout."""
        slide = self.new_slide("lead")
        self.text(slide, 0.93, 0.95, 8.4, 1.0, [title], size=title_size, color=TEAL_TITLE,
                  bold=True, anchor=MSO_ANCHOR.BOTTOM)
        self.text(slide, 0.93, 2.45, 7.8, 1.3, [subtitle], size=sub_size, color=WHITE, line_spacing=1.25)
        if extra:
            self.text(slide, 0.93, 4.95, 6.0, 0.35, [extra], size=11, color=MUTED_ON_PURPLE)
        if sticker and os.path.exists(sticker):
            slide.shapes.add_picture(sticker, Inches(7.55), Inches(3.35), height=Inches(1.9))
        return slide

    def stats(self, title, items, takeaway=None, subtitle=None):
        """Band slide with big-number cards: items = [(stat, body), ...]."""
        slide = self.band(title, subtitle)
        for (x, w), (stat, body) in zip(self.grid_x(len(items)), items):
            self.rect(slide, x, 2.35, w, 2.2, CARD)
            self.rect(slide, x, 2.35, w, 0.08, PURPLE)
            self.text(slide, x + 0.14, 2.62, w - 0.28, 0.55, [stat], size=30, color=TEAL, bold=True)
            self.text(slide, x + 0.14, 3.35, w - 0.28, 0.9, [self.runs(body)], size=11, color=DARK,
                      line_spacing=1.15)
        if takeaway:
            self.takeaway(slide, takeaway)
        return slide

    def cards(self, title, items, takeaway=None, subtitle=None, bar=PURPLE):
        """Band slide with 3-4 light cards: items = [(heading, body), ...]."""
        slide = self.band(title, subtitle)
        head_h = 0.55 if len(items) >= 4 else 0.3
        for (x, w), (head, body) in zip(self.grid_x(len(items)), items):
            self.card(slide, x, 2.3, w, 2.55, head, body, head_h=head_h, bar=bar)
        if takeaway:
            self.takeaway(slide, takeaway)
        return slide

    def tall_cards(self, title, items, takeaway=None):
        """Band slide with numbered tall cards overlapping the band:
        items = [(number, heading, body), ...]; 3 or 5 items look right."""
        slide = self.band(title)
        n = len(items)
        gap, h, num_size = (0.22, 3.5, 20) if n >= 5 else (0.25, 3.35, 22)
        for (x, w), (num, head, body) in zip(self.grid_x(n, gap=gap), items):
            self.tall_card(slide, x, 1.45, w, h, num, head, body, num_size=num_size)
        if takeaway:
            self.takeaway(slide, takeaway)
        return slide

    def panel(self, chip, title, items, note=None):
        """Left purple panel with logo; teal chip + title and a 2x2 grid of numbered
        cards on the right: items = [(number, lead, body), ...] (up to 4)."""
        slide = self.new_slide("panel")
        self.chip(slide, 4.2, 0.42, chip)
        self.text(slide, 4.2, 0.95, 5.55, 0.55, [title], size=20, color=PURPLE, bold=True)
        for i, (num, lead, body) in enumerate(items[:4]):
            px = 4.2 + (i % 2) * 2.85
            py = 1.70 + (i // 2) * 1.80
            self.panel_card(slide, px, py, 2.7, 1.65, num, lead, body)
        if note:
            # the layout's teal rule sits where the note goes - cover it first
            self.rect(slide, 4.1, 5.15, 5.9, 0.45, WHITE)
            self.text(slide, 4.2, 5.18, 5.5, 0.4, [note], size=9.5, color=PURPLE, anchor=MSO_ANCHOR.MIDDLE)
        return slide

    def code(self, title, code_lines, bullets):
        """Band slide: purple code panel left (comments teal, code white), teal-dot
        bullets right. bullets = markdown-ish strings (`code` spans are styled)."""
        slide = self.band(title)
        panel = self.rect(slide, CONTENT_X, 2.25, 4.45, 2.75, PURPLE, rounded=True, radius=0.045)
        tf = panel.text_frame
        tf.word_wrap = False
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left, tf.margin_right = Inches(0.22), Inches(0.12)
        tf.margin_top = tf.margin_bottom = Inches(0.12)
        first = True
        for line in code_lines:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.alignment = PP_ALIGN.LEFT
            p.line_spacing = 1.12
            if "#" in line and not line.strip().startswith("#"):
                code_part, comment = line.split("#", 1)
                parts = [(code_part, WHITE), ("#" + comment, TEAL)]
            elif line.strip().startswith("#"):
                parts = [(line, TEAL)]
            else:
                parts = [(line, WHITE)]
            for t, colr in parts:
                r = p.add_run()
                r.text = t
                r.font.name = MONO
                r.font.size = Pt(9.5)
                r.font.color.rgb = colr
        by = 2.4
        for b in bullets:
            self.oval(slide, 5.3, by + 0.04, 0.12, TEAL)
            self.text(slide, 5.55, by, 3.95, 0.7, [self.runs(b)], size=10.5, color=DARK, line_spacing=1.12)
            by += 0.85
        return slide

    def quote_panel(self, chip, title, lead, items, statement):
        """Quote layout: chip + title + italic lead + teal-oval bullet list on the
        white left half, a big teal split statement on the purple right panel.
        items = [(head, body), ...]; statement = [line1, line2]."""
        slide = self.new_slide("quote")
        self.rect(slide, 0, 0, 6.04, CANVAS_H, WHITE)   # hide the layout's quote furniture
        self.chip(slide, 0.45, 0.42, chip)
        self.text(slide, 0.45, 1.0, 5.3, 0.5, [title], size=22, color=PURPLE, bold=True)
        self.text(slide, 0.45, 1.58, 5.3, 0.4, [[{"t": lead, "i": True}]], size=10.5, color=GRAY,
                  line_spacing=1.2)
        self.rect(slide, 0.47, 2.18, 0.49, 0.05, TEAL)
        by = 2.5
        for head, body in items:
            self.oval(slide, 0.62, by + 0.05, 0.13, TEAL)
            self.text(slide, 0.92, by, 4.85, 0.5,
                      [[{"t": head, "b": True, "c": PURPLE, "s": 11.5}, {"t": "  " + body, "c": DARK, "s": 10.5}]],
                      line_spacing=1.12)
            by += 0.57
        self.text(slide, 6.5, 0.85, 3.1, 1.0, [statement[0]], size=24, color=TEAL_TITLE, bold=True, line_spacing=1.1)
        if len(statement) > 1:
            self.text(slide, 6.5, 2.25, 3.1, 2.4, [statement[1]], size=20, color=TEAL_TITLE, bold=True,
                      line_spacing=1.15)
        return slide

    def circle_trio(self, title, items, footer=None):
        """Purple full-bleed with three white circles carrying a flat teal glyph:
        items = [(glyph, label, quote), ...] where glyph is one of '?', 'bars', 'gear'
        or any single character."""
        slide = self.new_slide("team")
        self.text(slide, 2.0, 0.42, 6.0, 0.5, [title], size=22, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        cols = [(1.33, 0.77), (4.35, 3.80), (7.38, 6.82)]
        for (pic_x, col_x), (glyph, label, quote) in zip(cols, items):
            self.oval(slide, pic_x, 1.51, 1.3, WHITE)
            cx = pic_x + 0.65
            if glyph == "bars":
                for j in range(3):
                    self.rect(slide, cx - 0.3, 1.86 + j * 0.2, 0.6, 0.08, TEAL)
            elif glyph == "gear":
                g = slide.shapes.add_shape(MSO_SHAPE.GEAR_6, Inches(cx - 0.31), Inches(1.85), Inches(0.62), Inches(0.62))
                g.fill.solid()
                g.fill.fore_color.rgb = TEAL
                g.line.fill.background()
                g.shadow.inherit = False
            else:
                self.text(slide, cx - 0.4, 1.72, 0.8, 0.9, [glyph], size=44, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
            self.text(slide, col_x, 3.25, 2.4, 0.4, [label], size=14, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
            self.text(slide, col_x, 4.25, 2.4, 0.95, [quote], size=9.5, color=WHITE, align=PP_ALIGN.CENTER, line_spacing=1.15)
        if footer:
            self.text(slide, 1.5, 5.25, 7.0, 0.3, [footer], size=10, color=TEAL, align=PP_ALIGN.CENTER)
        return slide

    def closer(self, title="Thank you", name=None):
        """The template's own closing slide, placeholders kept and filled."""
        slide = self.new_slide("closer", keep_placeholders=True)
        phs = [shp for shp in slide.shapes if shp.has_text_frame]
        if phs:
            phs[0].text_frame.text = title
            if len(phs) > 1 and name:
                phs[1].text_frame.text = name
        return slide

    # ---- notes and output ---------------------------------------------------
    @staticmethod
    def notes(slide, text):
        slide.notes_slide.notes_text_frame.text = text

    def save(self, out, title=None, author=None):
        if title:
            self.prs.core_properties.title = title
        if author:
            self.prs.core_properties.author = author
        self.prs.save(out)
        n = len(self.prs.slides._sldIdLst)
        with_notes = sum(1 for s in self.prs.slides
                         if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip())
        print(f"saved {out}: {n} slides, {with_notes} with speaker notes, {os.path.getsize(out) // 1024} KB")
        return n


def marp_notes(marp_path):
    """Speaker-note blocks from a marp source, in slide order:
    <!--\\nSpeaker notes:\\n\\n...\\n-->  (the shape bob-2.0-refined.md uses)."""
    md = open(marp_path, encoding="utf-8").read()
    return re.findall(r"<!--\s*\nSpeaker notes:\s*\n\n(.*?)\n-->", md, re.DOTALL)
